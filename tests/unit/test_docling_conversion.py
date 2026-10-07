"""Real conversion orchestration and geometry with controlled Docling SDK output."""

import sys
import zipfile
from pathlib import Path
from types import ModuleType
from types import SimpleNamespace as NS
from unittest.mock import Mock

import pytest
from PIL import Image

from longparser.schemas import ProcessingConfig


class SDKObject:
    def __init__(self, **kwargs):
        self.__dict__.update(kwargs)


class OutputDocument:
    def __init__(self, items, pages=None):
        self.items = items
        self.pages = pages or {}
        self.texts = []

    def iterate_items(self, page_no=None):
        return iter((item, getattr(item, "level", 0)) for item in self.items + self.texts
                    if page_no is None or any(p.page_no == page_no for p in getattr(item, "prov", [])))


def bbox(left=50, top=50, right=150, bottom=150):
    box = NS(l=left, t=top, r=right, b=bottom, coord_origin="TOPLEFT")
    box.to_top_left_origin = lambda height: box
    return box


def item(text="x²", label="formula", page=1, box=None):
    return SDKObject(text=text, label=label, self_ref=f"#/texts/{text}",
                     prov=[NS(page_no=page, bbox=box or bbox())])


def page(image=None):
    return NS(size=NS(width=300, height=300), image=NS(pil_image=image or Image.new("RGB", (300, 300))))


@pytest.fixture
def converter(monkeypatch):
    from longparser.extractors import docling_extractor as module
    from longparser.extractors import latex_ocr
    sdk = Mock()
    factory = Mock(return_value=sdk)
    monkeypatch.setattr(module, "DocumentConverter", factory)
    for name in ["PdfPipelineOptions", "TesseractCliOcrOptions", "PdfFormatOption", "WordFormatOption",
                 "PowerpointFormatOption", "ExcelFormatOption", "CsvFormatOption"]:
        monkeypatch.setattr(module, name, SDKObject)
    monkeypatch.setattr(module, "InputFormat", NS(**{name: name for name in ["PDF", "DOCX", "PPTX", "XLSX", "CSV"]}))
    monkeypatch.setattr(module, "HierarchicalChunker", Mock(return_value=NS(chunk=lambda doc: [])))
    for name in ["TableItem", "PictureItem", "SectionHeaderItem", "TitleItem", "ListItem"]:
        monkeypatch.setattr(module, name, type(name, (SDKObject,), {}))
    monkeypatch.setattr(latex_ocr.MFDBackend, "_instance", NS(available=False))
    monkeypatch.setattr(latex_ocr.LaTeXOCR, "_instances", {})
    return NS(module=module, sdk=sdk, factory=factory, extractor=module.DoclingExtractor(), ocr=latex_ocr)


@pytest.mark.parametrize("mode,enabled", [("fast", False), ("smart", False), ("full", True)])
def test_converter_options_match_ocr_and_formula_configuration(converter, mode, enabled):
    converter.extractor._create_converter(ProcessingConfig(formula_mode=mode, do_ocr=False, export_images=True))
    options = converter.factory.call_args.kwargs["format_options"]
    assert set(options) == {"PDF", "DOCX", "PPTX", "XLSX", "CSV"}
    pipeline = options["PDF"].pipeline_options
    assert pipeline.do_ocr is False
    assert pipeline.do_formula_enrichment is enabled
    assert pipeline.generate_picture_images and pipeline.generate_page_images
    assert pipeline.ocr_options.tesseract_cmd == "tesseract"
    converter.extractor._create_converter(ProcessingConfig(formula_mode="full", formula_ocr=False))
    assert not converter.factory.call_args.kwargs["format_options"]["PDF"].pipeline_options.do_formula_enrichment
    converter.extractor._create_converter(ProcessingConfig(), formula_enrichment=True)
    assert converter.factory.call_args.kwargs["format_options"]["PDF"].pipeline_options.do_formula_enrichment


@pytest.mark.parametrize("mode", ["fast", "full", "smart"])
def test_pdf_modes_preserve_and_normalize_output(converter, mode):
    block = item(label="paragraph")
    doc = OutputDocument([block], {1: page()})
    converter.sdk.convert.return_value = NS(document=doc)
    result = converter.extractor._run_docling(Path("sample.pdf"), ProcessingConfig(formula_mode=mode))
    assert result.document is doc
    assert block.text == ("x²" if mode == "full" else "x^2")
    converter.extractor._run_docling(Path("sample.pdf"), ProcessingConfig(formula_mode=mode))
    assert converter.factory.call_count == 1


def test_huge_pdf_falls_back_and_sdk_failure_is_visible(converter):
    block = item()
    converter.sdk.convert.return_value = NS(document=OutputDocument([block], {n: page() for n in range(101)}))
    converter.extractor._run_docling(Path("huge.pdf"), ProcessingConfig(formula_mode="smart"))
    assert block.text == "x^2"
    converter.sdk.convert.side_effect = RuntimeError("converter failed")
    with pytest.raises(RuntimeError, match="converter failed"):
        converter.extractor._run_docling(Path("bad.pdf"), ProcessingConfig())


def test_docx_equations_use_external_parser_and_alignment_gate(converter, monkeypatch):
    sdk = ModuleType("docxlatex")
    sdk.Document = Mock(return_value=NS(get_equations=lambda: [r"\f r a c{x}{y}", " ", "z" * 1000]))
    monkeypatch.setitem(sys.modules, "docxlatex", sdk)
    blocks = [item(text="x/y"), item(text="abc")]
    converter.sdk.convert.return_value = NS(document=OutputDocument(blocks))
    converter.extractor._run_docling(Path("sample.docx"), ProcessingConfig(formula_mode="smart"))
    assert blocks[0].text == r"$$\frac{x}{y}$$"
    assert blocks[1].text == "abc"
    sdk.Document.side_effect = RuntimeError("broken document")
    assert converter.extractor._extract_docx_equations(Path("sample.docx")) == []
    monkeypatch.setitem(sys.modules, "docxlatex", None)
    assert converter.extractor._extract_docx_equations(Path("sample.docx")) == []


def test_pptx_equations_parse_real_xml_and_reject_broken_xml(converter, tmp_path):
    path = tmp_path / "math.pptx"
    xml = '<r xmlns:m="http://schemas.openxmlformats.org/officeDocument/2006/math"><m:oMath><m:t>x²</m:t><m:t>+ y</m:t></m:oMath></r>'
    with zipfile.ZipFile(path, "w") as archive:
        archive.writestr("ppt/slides/slide1.xml", xml)
        archive.writestr("ppt/slides/slide2.xml", "<broken")
    assert converter.extractor._extract_pptx_equations(path) == ["x² + y"]
    assert converter.extractor._extract_pptx_equations(tmp_path / "missing.pptx") == []
    block = item(text="x²")
    converter.sdk.convert.return_value = NS(document=OutputDocument([block]))
    converter.extractor._run_docling(path, ProcessingConfig(formula_mode="smart"))
    assert block.text == "$$x² + y$$"


def test_formula_merge_crop_and_missing_geometry(converter):
    first, second, third = item("a", box=bbox(50, 50, 150, 120)), item("b", box=bbox(50, 130, 150, 200)), item("c", box=bbox(220, 50, 290, 150))
    doc = OutputDocument([first, second, third], {1: page()})
    equations = converter.extractor._find_equation_items(doc)
    merged, unions, blanks = converter.extractor._merge_adjacent_formulas(equations, doc)
    assert merged == [(first, 1), (third, 1)]
    assert unions[id(first)] == (50, 50, 150, 200)
    assert blanks == {id(second)}
    crop = converter.extractor._crop_equation_bbox(doc, first, 1, unions)
    assert crop.size == (130, 195)
    assert converter.extractor._crop_equation_bbox(doc, first, 2) is None
    first.prov = []
    assert converter.extractor._crop_equation_bbox(doc, first, 1) is None
    first.prov = [NS(page_no=1, bbox=bbox(-1))]
    assert converter.extractor._crop_equation_bbox(doc, first, 1) is None
    first.prov = [NS(page_no=1, bbox=bbox(1, 1, 5, 5))]
    assert converter.extractor._crop_equation_bbox(doc, first, 1) is None
    assert converter.extractor._merge_adjacent_formulas([(first, 1)], doc) == ([(first, 1)], {}, set())


def test_smart_formula_ocr_merges_fragments_and_respects_limit(converter, monkeypatch):
    model = Mock(return_value="x+y")
    sdk = ModuleType("pix2tex.cli")
    sdk.LatexOCR = Mock(return_value=model)
    monkeypatch.setitem(sys.modules, "pix2tex.cli", sdk)
    first, fragment, later = item("a", box=bbox(50, 50, 150, 120)), item("b", box=bbox(50, 130, 150, 200)), item("c", page=2)
    doc = OutputDocument([first, fragment, later], {1: page(), 2: page()})
    converter.sdk.convert.return_value = NS(document=doc)
    converter.extractor._run_docling(Path("math.pdf"), ProcessingConfig(formula_mode="smart", smart_max_equations=1))
    assert first.text == "$$x+y$$"
    assert fragment.text == ""
    assert later.text == "c"


def test_math_page_scoring_validation_and_normalization(converter):
    doc = OutputDocument([item("αβγ", label="paragraph", page=2), item("plain", page=1), item("/C0", label="paragraph", page=3)])
    assert converter.extractor._detect_math_heavy_pages(doc) == [1, 2]
    assert converter.extractor._is_enriched_page_valid(doc, 1)
    assert not converter.extractor._is_enriched_page_valid(doc, 3)
    assert converter.extractor._is_enriched_page_valid(doc, 4)
    assert converter.extractor._normalize_unicode_math("$α$") == "$α$"
    assert converter.extractor._normalize_unicode_math("") == ""
    assert converter.extractor._normalize_latex("") == ""


def test_save_images_before_conversion_returns_empty(converter, tmp_path):
    assert converter.extractor.save_images(tmp_path) == []


def test_save_images_success_and_recoverable_failures(converter, tmp_path):
    picture = converter.module.PictureItem(self_ref="#/picture/1", get_image=lambda doc: Image.new("RGB", (20, 20)))
    table = converter.module.TableItem(self_ref="!!!", get_image=lambda doc: Image.new("RGB", (30, 30)))
    broken = converter.module.PictureItem(self_ref="broken", get_image=Mock(side_effect=ValueError("missing image")))
    converter.extractor._last_result = NS(document=OutputDocument([picture, table, broken], {1: page()}))
    saved = converter.extractor.save_images(tmp_path)
    assert len(saved) == 3
    assert {Image.open(path).size for path in saved} == {(300, 300), (20, 20), (30, 30)}
    assert all(path.suffix == ".png" and path.parent == tmp_path for path in saved)


def test_table_cells_dataframe_fallback_and_unavailable_data(converter):
    import pandas as pd
    extractor = converter.extractor
    table = converter.module.TableItem(data=NS(num_rows=1, num_cols=2, table_cells=[
        NS(start_row_offset_idx=0, end_row_offset_idx=1, start_col_offset_idx=0,
           end_col_offset_idx=2, text="merged")]))
    converted = extractor._build_table_from_item(table)
    assert (converted.n_rows, converted.n_cols, converted.cells[0].col_span) == (1, 2, 2)
    table.data.table_cells = [NS(text="broken")]
    table.export_to_dataframe = lambda doc: pd.DataFrame([["value", None]], columns=["Name", "Missing"])
    converted = extractor._build_table_from_item(table)
    assert [cell.text for cell in converted.cells] == ["Name", "Missing", "value", ""]
    table.export_to_dataframe = Mock(side_effect=ValueError("bad table"))
    assert extractor._build_table_from_item(table) is None
    table.data.num_rows = 0
    assert extractor._build_table_from_item(table) is None
    assert extractor._build_table_from_item(NS()) is None


def test_block_conversion_preserves_tables_equations_and_pptx_metadata(converter):
    module = converter.module
    child = item("table child", label="paragraph")
    table = module.TableItem(text="table content", label="table", self_ref="#/table/0", prov=[],
                            data=NS(num_rows=1, num_cols=1, table_cells=[NS(
                                start_row_offset_idx=0, end_row_offset_idx=1, start_col_offset_idx=0,
                                end_col_offset_idx=1, text="cell", ref=NS(cref=child.self_ref))]))
    equation = item("x+y")
    equation.export_to_markdown = Mock(side_effect=ValueError("markdown unavailable"))
    subtitle = item("Details", label="paragraph")
    footer = item("Footer", label="paragraph")
    list_item = module.ListItem(text="Step", label="list_item", self_ref="#/list/0", prov=[])
    doc = OutputDocument([child, table, equation, subtitle, footer, list_item,
                          item("1 / 22", label="paragraph"), item("x", label="paragraph"),
                          item("Header", label="page_header"), item("", label="paragraph")])
    def info(**kwargs):
        return module.PptxParaInfo(**dict(indent_level=0, is_title=False, is_subtitle=False, is_list=False, bullet_type="None") | kwargs)
    pptx = {0: {"Details": info(is_subtitle=True), "Footer": info(is_footer=True),
                "Step": info(indent_level=2)}}
    pages = converter.extractor._convert_to_pages(doc, {equation.self_ref: ["Math"]}, {}, {1: (300, 300)},
                                                  Path("sample.pptx"), "hash", pptx_text_map=pptx)
    blocks = [block for page in pages for block in page.blocks]
    assert not any(block.text in {"table child", "Footer", "1 / 22", "x", "Header"} for block in blocks)
    assert next(block for block in blocks if block.text == "Details").heading_level == 3
    assert next(block for block in blocks if block.text == "Step").indent_level == 2
    assert next(block for block in blocks if block.table).table.cells[0].text == "cell"
    formula = next(block for block in blocks if block.type == "equation")
    assert formula.hierarchy_path == ["Math"]
    assert formula.text == "⟦EQUATION⟧\nx+y\n⟦/EQUATION⟧"
    markdown = converter.extractor.to_markdown(NS(pages=pages))
    assert "### Details" in markdown and "    - Step" in markdown


def test_actual_pptx_indent_and_subtitle_map(converter, tmp_path):
    from pptx import Presentation
    from pptx.util import Inches
    prs = Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[0])
    slide.shapes.title.text = "Project title"
    slide.placeholders[1].text = "Subtitle"
    slide = prs.slides.add_slide(prs.slide_layouts[1])
    slide.shapes.title.text = "Details"
    body = slide.placeholders[1].text_frame
    body.paragraphs[0].text = "Main point"
    sub = body.add_paragraph()
    sub.text, sub.level = "Nested point", 2
    shape = slide.shapes.add_textbox(Inches(1), Inches(5), Inches(3), Inches(1))
    shape.text = "  Extra   text  "
    path = tmp_path / "slides.pptx"
    prs.save(path)
    mapping = converter.extractor._build_pptx_text_map(path)
    assert mapping[0]["Project title"].is_title
    assert mapping[0]["Subtitle"].is_subtitle
    assert mapping[1]["Nested point"].indent_level == 2
    assert mapping[1]["Nested point"].is_list
    assert "Extra text" in mapping[1]
    assert converter.extractor._build_pptx_text_map(tmp_path / "missing.pptx") == {}
    title = converter.module.TitleItem(text="Details", label="title", self_ref="#/title/0", prov=[])
    converter.sdk.convert.return_value = NS(document=OutputDocument([title, item("Nested point", label="paragraph")], {1: page()}))
    document, metadata = converter.extractor.extract(path, ProcessingConfig(formula_mode="fast"), page_numbers=[0])
    assert document.metadata.file_hash
    assert document.pages[0].blocks[0].heading_level == 2
    assert "PPTX mode" in metadata.reasons[0]


def test_pipeline_exports_real_document_chunks_and_manifests(converter, monkeypatch, tmp_path):
    import json

    from longparser.pipeline import orchestrator
    detector = ModuleType("fast_langdetect")
    detector.detect = lambda text: {"lang": "en", "score": 1.0}
    monkeypatch.setitem(sys.modules, "fast_langdetect", detector)
    path = tmp_path / "document.docx"
    path.write_bytes(b"controlled conversion boundary")
    paragraph = item("Return goods within thirty days of purchase.", label="paragraph")
    paragraph.prov = []
    doc = OutputDocument([paragraph], {1: page()})
    converter.sdk.convert.return_value = NS(document=doc)
    pipeline = orchestrator.PipelineOrchestrator()
    pipeline.extractor = converter.extractor
    result = pipeline.process_file(path, ProcessingConfig(formula_mode="fast", auto_detect_language=False))
    assert result.document.all_blocks[0].text == paragraph.text
    assert pipeline.chunk(result)[0].text == paragraph.text
    # The image exporter accepts a previously captured converter output.
    converter.extractor._last_result = converter.sdk.convert.return_value
    files = pipeline.export_results(result, tmp_path / "exports")
    blocks = json.loads(files["blocks"].read_text())
    manifest = json.loads(files["manifest"].read_text())
    assert blocks[0]["text"] == paragraph.text and "confidence" not in blocks[0]
    assert manifest["total_blocks"] == 1 and manifest["total_tables"] == 0
    assert files["markdown"].read_text().strip() == paragraph.text
    assert len(files["images"]) == 1
    assert json.loads(pipeline.export_chunks(result, tmp_path).read_text())[0]["text"] == paragraph.text
    assert json.loads(pipeline.export_hierarchy(result, tmp_path).read_text()) == []


def test_save_images_after_extract_uses_latest_converter_result(converter, tmp_path):
    path = tmp_path / "document.csv"
    path.write_text("name,value\nitem,1")
    converter.sdk.convert.return_value = NS(document=OutputDocument([item(label="paragraph")], {1: page()}))
    converter.extractor.extract(path, ProcessingConfig(formula_mode="fast"))
    assert len(converter.extractor.save_images(tmp_path / "images")) == 1


def test_smart_mfd_replaces_garbled_regions_without_duplicate_formulas(converter, monkeypatch):
    import numpy as np
    model = Mock(return_value="x+y")
    pix = ModuleType("pix2tex.cli")
    pix.LatexOCR = lambda: model
    monkeypatch.setitem(sys.modules, "pix2tex.cli", pix)
    detector = Mock()
    def detection(left, top, right, bottom, kind="isolated"):
        return {"box": np.array([[left, top], [right, top], [right, bottom], [left, bottom]]), "type": kind}
    detector.detect.return_value = [detection(50, 50, 150, 150), detection(160, 160, 260, 260),
                                   detection(1, 1, 10, 10)]
    mfd = converter.ocr.MFDBackend()
    mfd.available, mfd._mfd = True, detector
    monkeypatch.setattr(converter.ocr.MFDBackend, "_instance", mfd)
    formula = item("known", box=bbox(50, 50, 150, 150))
    garbled = item("α = x", label="paragraph", box=bbox(160, 160, 260, 260))
    formula.level = garbled.level = 1
    doc = OutputDocument([formula, garbled], {1: page()})
    converter.sdk.convert.return_value = NS(document=doc)
    converter.extractor._run_docling(Path("math.pdf"), ProcessingConfig(formula_mode="smart"))
    assert formula.text == "$$x+y$$"
    assert garbled.text == "$$x+y$$" and garbled.label == "formula"
    assert len(doc.items) == 2
    detector.detect.assert_called_once()


def test_smart_mfd_can_run_without_docling_formula_blocks(converter, monkeypatch):
    detector = Mock(return_value=[])
    monkeypatch.setattr(converter.ocr.MFDBackend, "_instance", NS(available=True, detect=detector))
    doc = OutputDocument([item("α = x", label="paragraph")], {1: page()})
    converter.sdk.convert.return_value = NS(document=doc)
    result = converter.extractor._run_docling(Path("math.pdf"), ProcessingConfig(formula_mode="smart"))
    assert result.document is doc


def test_extract_page_and_hierarchy_use_docling_provenance(converter, tmp_path):
    source = tmp_path / "data.csv"
    source.write_text("name,value\nitem,1")
    paragraph = item("Body", label="paragraph")
    converter.sdk.convert.return_value = NS(document=OutputDocument([paragraph], {1: page()}))
    config = ProcessingConfig(formula_mode="fast")
    assert converter.extractor.extract_page(source, 0, config).blocks[0].text == "Body"
    with pytest.raises(ValueError, match="Page 4 not found"):
        converter.extractor.extract_page(source, 4, config)
    converter.extractor._chunker = NS(chunk=lambda doc: [NS(text="Body", meta=NS(headings=["Section"], doc_items=[paragraph]))])
    hierarchy = converter.extractor.get_hierarchy(source, config)
    assert hierarchy[0].heading_path == ["Section"]
    assert hierarchy[0].page_number == 0 and hierarchy[0].order_index == 0


def test_equation_merge_skips_unrenderable_items_and_nonoverlapping_columns(converter):
    extractor = converter.extractor
    first = item("a", box=bbox(10, 20, 110, 120))
    second = item("b", box=bbox(180, 125, 280, 225))
    missing = item("c", page=2)
    doc = OutputDocument([first, second, missing], {1: page()})
    merged, unions, blanks = extractor._merge_adjacent_formulas([(first, 1), (second, 1), (missing, 2)], doc)
    assert merged == [(first, 1), (second, 1)]
    assert unions == {} and blanks == set()
    doc.pages[1].image = NS(pil_image=NS(size=None))
    assert extractor._crop_equation_bbox(doc, first, 1) is None
    assert extractor._merge_adjacent_formulas([(first, 1), (second, 1)], doc) == ([(first, 1), (second, 1)], {}, set())


def test_plaintext_markdown_escapes_leading_hash_and_preserves_running_header(converter):
    from longparser.schemas import BlockType
    from tests.conftest import make_block
    blocks = [make_block("# Python comment", BlockType.CODE), make_block("Running header", BlockType.HEADER)]
    assert converter.extractor.to_markdown(NS(pages=[NS(blocks=blocks)])) == "\\# Python comment\n\nRunning header\n"


def test_docling_sdk_metadata_fallbacks(converter):
    extractor = converter.extractor
    assert extractor._extract_bbox(NS(bbox=[1, 2, 3, 4])).model_dump() == {"x0": 1.0, "y0": 2.0, "x1": 3.0, "y1": 4.0}
    assert extractor._extract_bbox(NS(bbox="unknown")).x0 == 0
    assert extractor._extract_bbox(None).x1 == 0
    assert extractor._get_item_confidence(NS(confidence=.75)) == .75
    table = converter.module.TableItem(export_to_markdown=Mock(return_value="table"))
    assert extractor._get_item_text(table, "doc") == "table"
    table.export_to_markdown.assert_called_once_with(doc="doc")
    table.export_to_markdown.side_effect = ValueError("unrenderable")
    assert extractor._get_item_text(table) == ""
    fallback = NS(export_to_markdown=lambda: "fallback")
    assert extractor._get_item_text(fallback) == "fallback"
    assert not converter.module._is_mfd_candidate(1, [], 4)
