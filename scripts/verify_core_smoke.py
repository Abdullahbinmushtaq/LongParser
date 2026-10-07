"""Manually verify a non-editable core install with real DOCX extraction/chunking."""

from __future__ import annotations

import argparse
import json
import os
import socket
import sys
from importlib import metadata
from pathlib import Path


def main() -> None:
    """Run offline selected-path checks against an installed wheel.

    Args:
        None: Read the output directory and expected version from CLI arguments.

    Returns:
        None: Write JSON evidence and raise on an unmet assertion.
    """
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--output-dir", type=Path, required=True)
    parser.add_argument("--expected-version", default="0.1.5")
    args = parser.parse_args()
    args.output_dir.mkdir(parents=True, exist_ok=True)
    os.environ["HF_HUB_OFFLINE"] = "1"
    os.environ["TRANSFORMERS_OFFLINE"] = "1"
    connection_attempts = []

    def deny_connection(sock, address):
        connection_attempts.append(str(address))
        raise OSError("Network disabled for the installed-wheel smoke check")

    socket.socket.connect = deny_connection
    import longparser

    installed_path = Path(longparser.__file__).resolve()
    assert installed_path.is_relative_to(Path(sys.prefix).resolve()), installed_path
    assert longparser.__version__ == metadata.version("longparser") == args.expected_version
    assert not any(
        name.split(".")[0] in {"docling", "torch", "motor", "fastapi", "marker", "surya"}
        for name in sys.modules
    ), "Basic package import unexpectedly initialized a heavy dependency"
    distribution = metadata.distribution("longparser")
    direct_url = json.loads(distribution.read_text("direct_url.json") or "{}")
    assert not direct_url.get("dir_info", {}).get("editable", False), direct_url

    # Generate original fixture content; no downloaded/customer document is used.
    from docx import Document

    expected_texts = [
        "Unused items may be returned within thirty days. Keep the receipt as proof of purchase.",
        "Standard shipping arrives within five business days. A tracking number is included with each shipment.",
    ]
    fixture = args.output_dir / "smoke.docx"
    document = Document()
    document.add_heading("LongParser Smoke Test", 0)
    for heading, text in zip(("Return Policy", "Shipping Policy"), expected_texts, strict=True):
        document.add_heading(heading, 1)
        document.add_paragraph(text)
    document.save(fixture)

    from longparser import ChunkingConfig, DocumentPipeline, ProcessingConfig

    config = ProcessingConfig(
        backend="docling", do_ocr=False, formula_ocr=False, formula_mode="fast",
        auto_detect_language=False, export_images=False,
    )
    pipeline = DocumentPipeline(config)
    result = pipeline.process_file(fixture, config=config)
    assert result.chunks == [], "Extraction should precede explicit SDK chunking"
    blocks = result.document.all_blocks
    chunks = pipeline.chunk(result, ChunkingConfig(overlap_blocks=0))
    assert blocks and chunks and result.hierarchy
    assert chunks == result.chunks
    for text in expected_texts:
        assert text in "\n".join(block.text for block in blocks), text
        assert text in "\n".join(chunk.text for chunk in chunks), text
    block_ids = {block.block_id for block in blocks}
    pages_by_block = {block.block_id: page.page_number for page in result.document.pages
                      for block in page.blocks}
    for block in blocks:
        assert block.provenance.source_file == str(fixture)
        assert block.provenance.page_number == pages_by_block[block.block_id]
    for chunk in chunks:
        assert chunk.block_ids and set(chunk.block_ids) <= block_ids
        assert set(chunk.page_numbers) == {pages_by_block[bid] for bid in chunk.block_ids}
    assert not any(
        name.split(".")[0] in {"pymupdf4llm", "pymupdf", "fitz", "marker", "surya"}
        for name in sys.modules
    ), "Default DOCX processing loaded an optional GPL/AGPL backend"
    evidence = {
        "version": longparser.__version__, "python": sys.version.split()[0],
        "installed_package": str(installed_path), "editable_install": False,
        "fixture": str(fixture), "backend": "docling", "ocr": False,
        "pages": len(result.document.pages), "blocks": len(blocks),
        "page_numbers": [page.page_number for page in result.document.pages],
        "hierarchy_items": len(result.hierarchy), "chunks": len(chunks),
        "expected_texts_preserved": True, "source_references_valid": True,
        "optional_backend_imports": [], "blocked_connections": connection_attempts,
        "docling": metadata.version("docling"),
        "chunk_texts": [chunk.text for chunk in chunks],
    }
    serialized = json.dumps(evidence, indent=2) + "\n"
    (args.output_dir / "smoke-result.json").write_text(serialized)
    print(serialized)


if __name__ == "__main__":
    main()
